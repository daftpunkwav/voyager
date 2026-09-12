"""Doc submodule tests: import/parse pipeline (injected parse_fn), section
reads, validation, removal events.

No real PDF/docx parsing: the worker gets parse_fn injected; extractor
pure functions are unit-tested locally (txt/epub/docx built and parsed in
the test, PDF only covers the rejection path).
"""

import asyncio
import zipfile
from pathlib import Path

import pytest
from platform_actor import ActorContext
from platform_capability import execute
from platform_contracts import LOCAL_USER, ActorKind, ActorRef, ServiceError
from platform_eventbus import EventBus, EventLog
from platform_secrets import SecretStore
from sources.capabilities import SourcesDeps, init_all, registry
from sources.modules.doc.extract import (
    ExtractError,
    _from_docx,
    _from_text,
    extract_sections,
)
from sources.modules.doc.store import DocStore
from sources.modules.doc.worker import DocWorker
from sources.modules.repo.store import RepoStore
from sources.modules.web.store import WebStore

USER_CTX = ActorContext(actor=LOCAL_USER)
AGENT_CTX = ActorContext(actor=ActorRef(kind=ActorKind.AGENT, id="agent.main", scopes=()))


@pytest.fixture()
def deps(tmp_path):
    log = EventLog(tmp_path / "events.db")
    d = SourcesDeps(
        repo_store=RepoStore(tmp_path / "repo.db"),
        doc_store=DocStore(tmp_path / "doc.db"),
        web_store=WebStore(tmp_path / "web.db"),
        secrets=SecretStore(tmp_path / "secrets.db", key_material="t"),
        bus=EventBus(log),
        repo_queue=asyncio.Queue(),
        doc_queue=asyncio.Queue(),
        workspace=tmp_path / "ws",
    )
    init_all(d)
    yield d, log
    d.repo_store.close()
    d.doc_store.close()
    d.web_store.close()
    log.close()


def _seed_file(ws: Path, name: str, content: bytes | str) -> Path:
    ws.mkdir(parents=True, exist_ok=True)
    p = ws / name
    if isinstance(content, str):
        p.write_text(content, encoding="utf-8")
    else:
        p.write_bytes(content)
    return p


class TestAddDocument:
    async def test_add_parseable_enqueues_and_emits_added(self, deps, tmp_path) -> None:
        d, log = deps
        src = _seed_file(tmp_path / "ws", "report.pdf", b"%PDF-1.4 fake")
        ref = await execute(
            registry,
            "add_document",
            USER_CTX,
            {"file_path": str(src), "title": "quarterly report", "tags": ["work"]},
        )
        doc = d.doc_store.get(ref.job_id)
        assert doc["status"] == "parsing"  # parseable formats are enqueued for the worker
        assert doc["ext"] == ".pdf"
        assert d.doc_queue.qsize() == 1
        assert Path(doc["local_path"]).parent == Path(d.workspace) / "doc"
        types = [e.type for _, e in log.read_after()]
        assert "source.added" in types

    async def test_add_unknown_ext_stored_without_parse(self, deps, tmp_path) -> None:
        """Unknown extension = archive semantics: stored as-is, never enqueued
        for parsing.
        """
        d, _ = deps
        src = _seed_file(tmp_path / "ws", "dataset.zip", b"PK fake")
        ref = await execute(registry, "add_document", USER_CTX, {"file_path": str(src)})
        doc = d.doc_store.get(ref.job_id)
        assert doc["status"] == "stored"
        assert d.doc_queue.qsize() == 0

    async def test_missing_file_and_outside_workspace(self, deps, tmp_path) -> None:
        with pytest.raises(ServiceError, match="File not found"):
            await execute(
                registry, "add_document", USER_CTX, {"file_path": str(tmp_path / "ws" / "nope.pdf")}
            )
        # paths outside the jail are not disclosed, existing or not
        with pytest.raises(ServiceError, match="workspace"):
            await execute(
                registry, "add_document", USER_CTX, {"file_path": str(tmp_path / "nope.pdf")}
            )
        outside = tmp_path / "outside"
        outside.mkdir()
        src = outside / "a.pdf"
        src.write_bytes(b"x")
        with pytest.raises(ServiceError, match="workspace"):
            await execute(registry, "add_document", USER_CTX, {"file_path": str(src)})

    async def test_title_filename_sanitized(self, deps, tmp_path) -> None:
        """A title with path separators/reserved characters must not escape
        workspace/doc.
        """
        d, _ = deps
        src = _seed_file(tmp_path / "ws", "a.md", b"# t")
        ref = await execute(
            registry,
            "add_document",
            USER_CTX,
            {"file_path": str(src), "title": '../../evil "name"'},
        )
        local = Path(d.doc_store.get(ref.job_id)["local_path"])
        assert local.parent == Path(d.workspace) / "doc"
        assert local.is_file()


class TestParseWorker:
    async def test_parse_then_ready_with_sections(self, deps, tmp_path) -> None:
        d, log = deps

        def fake_parse(path: Path, ext: str):
            return _from_text(Path(path))  # reuse the real text chunking

        worker = DocWorker(
            d.doc_store, EventBus(log), d.doc_queue, tmp_path / "ws", parse_fn=fake_parse
        )
        await worker.start()
        src = _seed_file(tmp_path / "ws", "book.md", "# chapter\n" + "body" * 50)
        ref = await execute(
            registry, "add_document", USER_CTX, {"file_path": str(src), "title": "handbook"}
        )  # enqueued internally
        # Poll for the terminal EVENT, not the status: the worker flips the
        # status first and emits source.ready several awaits later, and stop()
        # cancels the task — polling the status races the emit under load.
        for _ in range(200):
            await asyncio.sleep(0.02)
            if log.read_after(types=["source.ready"]):
                break
        await worker.stop()
        doc = d.doc_store.get(ref.job_id)
        assert doc["status"] == "ready"
        outline = d.doc_store.sections_outline(ref.job_id)
        assert len(outline) >= 1
        section = d.doc_store.section(ref.job_id, outline[0]["section_no"])
        assert section["text"].startswith("# chapter")
        types = [e.type for _, e in log.read_after(types=["source.ready", "task.progress"])]
        assert "source.ready" in types

    async def test_parse_failure_marks_failed(self, deps, tmp_path) -> None:
        d, log = deps

        def boom(path, ext):
            raise ExtractError("Failed to open PDF: test injection")

        worker = DocWorker(d.doc_store, EventBus(log), d.doc_queue, tmp_path / "ws", parse_fn=boom)
        await worker.start()
        src = _seed_file(tmp_path / "ws", "bad.pdf", b"%PDF broken")
        ref = await execute(
            registry, "add_document", USER_CTX, {"file_path": str(src)}
        )  # enqueued internally
        # Same event-first polling as above: task.failed lands after the
        # status flip; stop() would cancel an in-flight emit.
        for _ in range(200):
            await asyncio.sleep(0.02)
            if log.read_after(types=["task.failed"]):
                break
        await worker.stop()
        doc = d.doc_store.get(ref.job_id)
        assert doc["status"] == "failed" and "PDF" in doc["error"]
        types = [e.type for _, e in log.read_after(types=["task.failed"])]
        assert types == ["task.failed"]

    async def test_remove_cleans_record_and_file(self, deps, tmp_path) -> None:
        d, log = deps

        def noop(path, ext):
            return []

        worker = DocWorker(d.doc_store, EventBus(log), d.doc_queue, tmp_path / "ws", parse_fn=noop)
        await worker.start()
        src = _seed_file(tmp_path / "ws", "x.md", b"hi")
        ref = await execute(registry, "add_document", USER_CTX, {"file_path": str(src)})
        local = d.doc_store.get(ref.job_id)["local_path"]
        await execute(registry, "remove_document", USER_CTX, {"doc_id": ref.job_id})
        # The worker deletes the file asynchronously; wait for the deletion
        # itself instead of a fixed sleep (slow under a loaded session).
        for _ in range(200):
            await asyncio.sleep(0.02)
            if not Path(local).exists():
                break
        await worker.stop()
        assert d.doc_store.get(ref.job_id) is None
        assert not Path(local).exists()
        types = [e.type for _, e in log.read_after(types=["source.removed"])]
        assert types == ["source.removed"]


class TestReadAndMeta:
    async def test_get_document_outline_and_section(self, deps, tmp_path) -> None:
        d, _ = deps
        src = _seed_file(tmp_path / "ws", "r.md", b"body text")
        ref = await execute(
            registry, "add_document", USER_CTX, {"file_path": str(src), "title": "report"}
        )
        d.doc_store.replace_sections(
            ref.job_id,
            [
                {
                    "section_no": 1,
                    "title": "Overview",
                    "page_start": 1,
                    "page_end": 2,
                    "text": "chapter one text",
                },
                {
                    "section_no": 2,
                    "title": "Method",
                    "page_start": 3,
                    "page_end": 9,
                    "text": "chapter two text",
                },
            ],
        )
        d.doc_store.set_status(ref.job_id, "ready")
        detail = await execute(registry, "get_document", AGENT_CTX, {"doc_id": ref.job_id})
        assert [s["title"] for s in detail["sections"]] == ["Overview", "Method"]
        assert "text" not in detail["sections"][0]  # outline carries no body text
        section = await execute(
            registry, "get_doc_section", AGENT_CTX, {"doc_id": ref.job_id, "section_no": 2}
        )
        assert section["text"] == "chapter two text"
        assert section["total_sections"] == 2
        with pytest.raises(ServiceError, match="Section not found"):
            await execute(
                registry, "get_doc_section", AGENT_CTX, {"doc_id": ref.job_id, "section_no": 9}
            )

    async def test_search_sections_with_snippet(self, deps, tmp_path) -> None:
        d, _ = deps
        src = _seed_file(tmp_path / "ws", "r.md", b"x")
        ref = await execute(
            registry, "add_document", USER_CTX, {"file_path": str(src), "title": "AI handbook"}
        )
        d.doc_store.replace_sections(
            ref.job_id,
            [
                {
                    "section_no": 1,
                    "title": "",
                    "page_start": 0,
                    "page_end": 0,
                    "text": "前面铺垫。" * 20 + "注意力机制是核心。" + "后面内容。" * 20,
                }
            ],
        )
        hits = await execute(registry, "search_documents", AGENT_CTX, {"query": "注意力机制"})
        assert hits[0]["doc_id"] == ref.job_id and hits[0]["section_no"] == 1
        assert "注意力机制" in hits[0]["snippet"]

    async def test_set_meta_and_tag_validation(self, deps, tmp_path) -> None:
        d, _ = deps
        src = _seed_file(tmp_path / "ws", "m.md", b"x")
        ref = await execute(
            registry, "add_document", USER_CTX, {"file_path": str(src), "title": "notes collection"}
        )
        await execute(
            registry,
            "set_document_meta",
            USER_CTX,
            {
                "doc_id": ref.job_id,
                "category": "courses",
                "tags": ["ml"],
                "progress": "learning",
                "note": "reading",
            },
        )
        doc = d.doc_store.get(ref.job_id)
        assert doc["category"] == "courses" and doc["tags"] == ["ml"]
        with pytest.raises(ServiceError, match="Invalid tag"):
            await execute(
                registry, "set_document_meta", USER_CTX, {"doc_id": ref.job_id, "tags": ['a"b']}
            )

    async def test_list_filter_by_status_and_query(self, deps, tmp_path) -> None:
        d, _ = deps
        s1 = _seed_file(tmp_path / "ws", "a.md", b"x")
        s2 = _seed_file(tmp_path / "ws", "b.zip", b"z")
        r1 = await execute(
            registry, "add_document", USER_CTX, {"file_path": str(s1), "title": "Alpha handbook"}
        )
        await execute(
            registry, "add_document", USER_CTX, {"file_path": str(s2), "title": "Beta dataset"}
        )
        d.doc_store.set_status(r1.job_id, "ready")
        ready_only = await execute(registry, "list_documents", USER_CTX, {"status": "ready"})
        assert [r["title"] for r in ready_only] == ["Alpha handbook"]
        by_query = await execute(registry, "list_documents", USER_CTX, {"query": "alpha"})
        assert len(by_query) == 1


class TestExtractor:
    def test_text_chunks_long_content(self, tmp_path) -> None:
        p = tmp_path / "long.md"
        p.write_text(("paragraph text.\n\n" * 2000), encoding="utf-8")
        sections = extract_sections(p, ".md")
        assert len(sections) > 1  # long text is chunked on blank lines into chapters
        assert [s.section_no for s in sections] == list(range(1, len(sections) + 1))

    def test_docx_heading_split(self, tmp_path) -> None:
        import docx

        document = docx.Document()
        document.add_paragraph("Chapter 1", style="Heading 1")
        document.add_paragraph("Chapter 1 body")
        document.add_paragraph("Chapter 2", style="Heading 1")
        document.add_paragraph("Chapter 2 body")
        p = tmp_path / "d.docx"
        document.save(str(p))
        sections = _from_docx(p)
        assert [s.title for s in sections] == ["Chapter 1", "Chapter 2"]

    def test_epub_spine_order(self, tmp_path) -> None:
        p = tmp_path / "b.epub"
        with zipfile.ZipFile(p, "w") as zf:
            zf.writestr("mimetype", "application/epub+zip")
            zf.writestr("content/ch1.xhtml", "<html><body><h1>One</h1><p>A</p></body></html>")
            zf.writestr("content/ch2.xhtml", "<html><body><h1>Two</h1><p>B</p></body></html>")
            zf.writestr(
                "content/book.opf",
                "<package><manifest>"
                '<item id="c1" href="ch1.xhtml"/>'
                '<item id="c2" href="ch2.xhtml"/>'
                "</manifest><spine>"
                '<itemref idref="c1"/><itemref idref="c2"/>'
                "</spine></package>",
            )
        sections = extract_sections(p, ".epub")
        assert [s.title for s in sections] == ["One", "Two"]

    def test_pdf_corrupt_raises_extract_error(self, tmp_path) -> None:
        p = tmp_path / "bad.pdf"
        p.write_bytes(b"not a pdf at all")
        with pytest.raises(ExtractError):
            extract_sections(p, ".pdf")

    def test_empty_text_raises(self, tmp_path) -> None:
        p = tmp_path / "e.txt"
        p.write_text("  \n  ", encoding="utf-8")
        with pytest.raises(ExtractError, match="empty"):
            extract_sections(p, ".txt")
